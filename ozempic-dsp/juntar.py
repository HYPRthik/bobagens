#!/usr/bin/env python3
"""
Resolve as lojas do Ozempic em enderecos.

A planilha traz CNPJ, Razao Social, Nome Fantasia, Municipio, UF e coordenada —
mas NENHUM endereco de rua, apesar do "enderecos" no nome do arquivo. O
geofencing da Yahoo le um endereco por linha em texto livre, entao coordenada
nao serve.

Duas saidas, nesta ordem:
  1. Cruzar a coordenada com as bases de endereco ja processadas nesta pasta.
     So aceita quando a distancia e pequena E os nomes compartilham um termo
     proprio — a menos de 60 m de uma farmacia costuma haver outra farmacia, e
     colar o endereco do vizinho e pior do que nao ter endereco.
  2. O que sobra vai pelo nome do POI (--nome-sem-endereco), com expectativa
     baixa: a DSP so aceita POI pelo nome em Airports, Arena/Stadiums e
     Universities/Colleges.

Uso:  python3 juntar.py
"""
import csv
import io
import math
import os
import re
import sys
import unicodedata

from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(BASE)
XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "origem", "ozempic.xlsx")
DIST_MAX = 60.0        # metros
BASES = ["farmacias-dsp/origem_farmacias.csv", "beleza-dsp/origem_beleza.csv",
         "cinemas-dsp/origem_cinemas.csv", "pois-jlr-dsp/origem_pois_completa.csv",
         "mercadolivre-ooh-dsp/origem_ml_ooh.csv",
         "campari-negroni-dsp/origem/pdvs_base_completa.csv"]
# Termos genericos do ramo e da razao social: nao provam que e a mesma loja.
# Marca NAO entra aqui — "Drogasil" identifica a loja, "drogaria" nao.
GENERICOS = {"farmacia", "farmacias", "drogaria", "drogarias", "droga", "drogas",
             "ltda", "eireli", "cia", "comercio", "produtos", "farmaceuticos",
             "loja", "lojas", "filial", "matriz"}


def sem_acento(s):
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()


def tokens(s):
    return {t for t in re.findall(r"[a-z]{3,}", sem_acento(s).lower())}


def proprios(s):
    """Termos que identificam a loja, tirando o que e generico do ramo."""
    return tokens(s) - GENERICOS


def hav(a, b, c, d):
    R = 6371000.0
    p1, p2 = math.radians(a), math.radians(c)
    x = (math.sin((p2 - p1) / 2) ** 2
         + math.cos(p1) * math.cos(p2) * math.sin(math.radians(d - b) / 2) ** 2)
    return 2 * R * math.asin(math.sqrt(x))


def val(x):
    if x is None:
        return ""
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x).strip()


# ------------------------------------------------------------------ leitura
wb = load_workbook(XLSX, data_only=True)
rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
wb.close()
# a planilha tem titulo e fonte antes do cabecalho, e TOTAL/rodape depois
h = next(n for n, r in enumerate(rows) if val(r[0]) == "CNPJ")
hdr = [val(x) for x in rows[h]]
col = {x: n for n, x in enumerate(hdr) if x}
lojas = [r for r in rows[h + 1:] if val(r[col["CNPJ"]]).isdigit()]
print(f"lojas na planilha: {len(lojas)}  (cabecalho na linha {h + 1}; TOTAL e rodape fora)")

# ------------------------------------------------------------------ indice
pts = []
for b in BASES:
    p = os.path.join(RAIZ, b)
    if not os.path.exists(p):
        continue
    for r in csv.DictReader(io.StringIO(open(p, encoding="utf-8-sig").read())):
        try:
            pts.append((float(r["lat"]), float(r["lng"]), r, b.split("/")[0]))
        except (ValueError, TypeError, KeyError):
            pass
print(f"pontos com endereco nas bases: {len(pts)}")

# ------------------------------------------------------------------ cruzamento
CAMPOS = ["nome", "endereco", "cidade", "uf", "cep", "lat", "lng", "cnpj", "notas", "origem_match"]
casados, semend, rejeitados = [], [], []
for r in lojas:
    nome = val(r[col["Nome Fantasia"]]) or val(r[col["Razão Social"]])
    reg = {"nome": nome, "endereco": "", "cidade": val(r[col["Município"]]),
           "uf": val(r[col["UF"]]), "cep": "", "lat": val(r[col["Latitude"]]),
           "lng": val(r[col["Longitude"]]), "cnpj": val(r[col["CNPJ"]]),
           "notas": val(r[col["Notas Fiscais"]]), "origem_match": ""}
    try:
        la, lo = float(reg["lat"]), float(reg["lng"])
    except ValueError:
        semend.append(reg)
        continue
    cand = [(hav(la, lo, a, b), m, f) for a, b, m, f in pts if abs(a - la) < 0.002 and abs(b - lo) < 0.002]
    cand.sort(key=lambda t: t[0])
    achou = None
    for dist, m, fonte in cand:
        if dist > DIST_MAX:
            break
        # exige termo proprio em comum: perto de farmacia ha outra farmacia
        if proprios(nome) & proprios(m.get("name", "") or m.get("storechain", "")):
            achou = (dist, m, fonte)
            break
        rejeitados.append((nome, dist, m.get("name", ""), m["formatted_address"][:60]))
    if achou:
        dist, m, fonte = achou
        reg.update(endereco=m["formatted_address"], cidade=m.get("city") or reg["cidade"],
                   uf=m.get("uf") or reg["uf"], cep=m.get("cep", ""),
                   origem_match=f"{fonte} ({dist:.0f} m)")
        casados.append(reg)
    else:
        semend.append(reg)

os.makedirs(os.path.join(BASE, "enriquecido"), exist_ok=True)
with open(os.path.join(BASE, "enriquecido", "ozempic.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=CAMPOS)
    w.writeheader()
    w.writerows(casados + semend)
with open(os.path.join(BASE, "enriquecido", "ozempic_SEM_ENDERECO.csv"), "w",
          newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";")
    w.writerow(["CNPJ", "Nome Fantasia", "Municipio", "UF", "Latitude", "Longitude",
                "Notas Fiscais", "ENDERECO_A_PREENCHER"])
    for g in semend:
        w.writerow([g["cnpj"], g["nome"], g["cidade"], g["uf"], g["lat"], g["lng"],
                    g["notas"], ""])

print(f"\nendereco recuperado das bases : {len(casados)}")
for g in casados:
    print(f"   {g['nome'][:30]:32} {g['origem_match']:26} {g['endereco'][:46]}")
print(f"sem endereco                  : {len(semend)}")
print(f"\nvizinhos rejeitados (nome nao bate): {len(rejeitados)}")
for n, d, mn, e in rejeitados[:6]:
    print(f"   {n[:26]:28} x {mn[:20]:22} {d:4.0f} m  {e[:40]}")
